"""
Lecture d'un document importé (Word, PDF numérique, PDF scanné) en un seul
« document propre » : titres, paragraphes, gras, italique, tableaux Word.

Le même résultat sert à l'écran « Analyse des risques » (html) et à l'import de
modèle et à la contrathèque (texte : un paragraphe par bloc, séparés par une
ligne vide).

Principes :
- un document propre garde sa mise en forme : le texte n'est jamais « corrigé »
  (l'ancienne correction d'espaces abîmait montants, dates, articles de loi et
  adresses : « 12. 500, 00 € », « L. 1242-8 », « www. lumenjuris. com ») ;
- on retire ce qui n'appartient pas au contrat : en-têtes, pieds de page,
  numéros de page ;
- on recolle les mots coupés en fin de ligne (en gardant le tiret des mots
  composés) et les paragraphes coupés par un changement de page ;
- un scan est lu EN ENTIER par Google Vision, par lots de 5 pages (limite de
  l'API, qui ignorait sans le dire les pages suivantes), en gardant les
  paragraphes reconnus par Vision ;
- la décision « scan ou numérique » se prend page par page : un contrat
  numérique suivi d'une page de signature scannée est lu correctement.
"""
from __future__ import annotations

import html as html_module
import io
import logging
import math
import re
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from typing import List, Optional, Sequence, Set, Tuple

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import mammoth
except ImportError:
    mammoth = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from google.cloud import vision
except ImportError:
    vision = None

from .pdf_processing import vision_client

logger = logging.getLogger(__name__)

MESSAGE_DOC = (
    "Ancien format Word (.doc) non pris en charge : ouvrez le fichier dans Word, choisissez "
    "« Enregistrer sous » au format .docx (ou PDF), puis importez-le à nouveau."
)
MESSAGE_VIDE = "Aucun texte lisible dans ce document."


@dataclass
class DocumentLu:
    texte: str
    html: Optional[str]
    methode: str  # "word" ou "server" (valeurs déjà attendues par le front)
    lecture: str  # "word", "numerique", "ocr" ou "mixte"
    pages: int


@dataclass
class Bloc:
    tag: str  # "h1", "h2", "bold" (titre de clause, rendu en <p>) ou "p"
    texte: str
    html: str  # contenu en ligne (gras, italique), sans la balise du bloc
    page: int = 0


# ─── Caractères ────────────────────────────────────────────────────────────────

# Espaces insécables, tirets « spéciaux » et ligatures des PDF : ils empêchent de
# retrouver une clause dans le texte (l'IA renvoie des caractères ordinaires).
_REMPLACEMENTS = str.maketrans({
    " ": " ", " ": " ", " ": " ", " ": " ", " ": " ",
    "‐": "-", "‑": "-",
    "ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl", "ﬅ": "st", "ﬆ": "st",
    "​": None, "‌": None, "‍": None, "﻿": None,
})


def normaliser(texte: str) -> str:
    """Remplace les caractères invisibles ou « spéciaux » ; le tiret conditionnel (U+00AD) est gardé pour les césures."""
    return texte.translate(_REMPLACEMENTS)


def espaces_ocr(texte: str) -> str:
    """Ajoute l'espace oubliée par la lecture d'un scan entre deux phrases (« fin.Le ») ou après une virgule entre deux mots.

    Volontairement restreint aux lettres : les chiffres (12.500,00), sigles (S.A.S.), adresses
    (www.lumenjuris.com) et références (L.1242-8) ne sont jamais touchés.
    """
    # Guillemets français lus comme deux chevrons (« << le Prestataire >> »).
    texte = texte.replace("<<", "«").replace(">>", "»")
    texte = re.sub(r"(?<=[a-zà-öø-ÿ]{2})([.;!?])(?=[A-ZÀ-ÖØ-Þ][a-zà-öø-ÿ])", r"\1 ", texte)
    texte = re.sub(r"(?<=[a-zà-öø-ÿ]),(?=[a-zà-öø-ÿ]{2})", ", ", texte)
    return re.sub(r" {2,}", " ", texte)


# ─── Mots coupés en fin de ligne ───────────────────────────────────────────────

_LETTRES = "a-zA-Zà-öø-ÿÀ-ÖØ-ÞœŒæÆ"
_FIN_COUPEE = re.compile(rf"([{_LETTRES}]+)[-­]$")
_DEBUT_MINUSCULE = re.compile(r"^([a-zà-öø-ÿœæ]+)")
_TIRET_FINAL_HTML = re.compile(r"[-­]((?:</(?:strong|em)>)*)\s*$")

# Mots composés usuels : coupés sur leur tiret en fin de ligne, ils le gardent.
_DEBUTS_COMPOSES = {
    "sous", "vice", "ci", "non", "ex", "co", "auto", "anti", "avant", "après", "arrière", "contre", "demi", "semi",
    "mi", "porte", "grand", "haut", "bas", "nord", "sud", "est", "ouest", "pré", "post", "inter", "extra", "multi",
    "quasi", "pseudo", "néo", "bien", "mal", "sans", "au", "là", "vis", "peut", "c'est", "est-ce", "dix", "vingt",
    "trente", "quarante", "cinquante", "soixante", "quatre", "franco", "lieu", "chef", "procès", "mise", "plus",
    "moins", "état", "outre", "court", "long", "moyen", "tiers", "garde", "beau", "belle", "petit", "petite",
}
_FINS_COMPOSEES = {
    "ci", "là", "même", "mêmes", "dire", "être", "ce", "il", "elle", "on", "je", "tu", "nous", "vous", "ils",
    "elles", "t", "y", "en", "dessus", "dessous", "delà", "deçà", "après", "avant", "contre",
}


def vocabulaire(texte: str) -> Set[str]:
    """Mots du document (y compris les mots composés), pour trancher entre césure et mot composé."""
    return {m.lower() for m in re.findall(rf"[{_LETTRES}]+(?:-[{_LETTRES}]+)*", texte)}


def mode_cesure(avant: str, apres: str, mots: Set[str]) -> Optional[str]:
    """Raccord d'une ligne finissant par un tiret avec la suivante.

    None : pas une coupure de mot (on recolle avec une espace) ; "" : coupure typographique
    (« résilia- / tion » → « résiliation ») ; "-" : mot composé (« vice- / président »).
    """
    fin, debut = _FIN_COUPEE.search(avant), _DEBUT_MINUSCULE.match(apres)
    if not fin or not debut:
        return None
    if avant.endswith("­"):
        return ""
    a, b = fin.group(1).lower(), debut.group(1)
    if a + b in mots:
        return ""
    if f"{a}-{b}" in mots or a in _DEBUTS_COMPOSES or b in _FINS_COMPOSEES:
        return "-"
    return ""


def _l_lu_comme_i(avant: str, t: str, h: str, mots: Set[str]) -> Tuple[str, str]:
    """Scan : « difficu- / Ité » — le « l » en début de ligne lu « I ». Corrigé seulement si le mot existe dans le document."""
    fin, suite = _FIN_COUPEE.search(avant), re.match(r"^I([a-zà-öø-ÿ]+)", t)
    if fin and suite and (fin.group(1) + "l" + suite.group(1)).lower() in mots and h.startswith("I"):
        return "l" + t[1:], "l" + h[1:]
    return t, h


def joindre(lignes: Sequence[Tuple[str, str]], mots: Set[str]) -> Tuple[str, str]:
    """Recolle des lignes (texte, html) en un paragraphe, en réparant les mots coupés."""
    texte, html = "", ""
    for t, h in lignes:
        t, h = t.strip(), h.strip()
        if not t:
            continue
        if not texte:
            texte, html = t, h
            continue
        t, h = _l_lu_comme_i(texte, t, h, mots)
        mode = mode_cesure(texte, t, mots)
        if mode is None:
            texte, html = f"{texte} {t}", f"{html} {h}"
        else:
            texte = texte[:-1] + mode + t
            html = _TIRET_FINAL_HTML.sub(lambda m: mode + m.group(1), html, count=1) + h
    return texte.replace("­", ""), html.replace("­", "")


# ─── En-têtes, pieds de page, numéros de page ──────────────────────────────────

_NUMERO_PAGE = re.compile(r"^(?:page|p\.)?\s*#(?:\s*(?:/|sur|of|de)\s*#)?$|^-\s*#\s*-$", re.I)
_MARGE_HAUTE, _MARGE_BASSE = 0.09, 0.91


def _cle(texte: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\d+", "#", texte.lower())).strip()


def parasites(elements: Sequence[Tuple[int, float, float, str]], nb_pages: int) -> Set[int]:
    """Indices des éléments (page, haut, bas, texte) qui sont des en-têtes, pieds ou numéros de page.

    haut / bas : position verticale en fraction de la hauteur de page. Un élément n'est retiré
    que s'il est dans la marge ET qu'il est un numéro de page ou se répète sur au moins la
    moitié des pages : un titre en haut de la première page n'est jamais retiré.
    """
    en_marge = [i for i, (_, haut, bas, _) in enumerate(elements) if haut < _MARGE_HAUTE or bas > _MARGE_BASSE]
    pages_par_cle = {}
    for i in en_marge:
        pages_par_cle.setdefault(_cle(elements[i][3]), set()).add(elements[i][0])
    seuil = max(2, math.ceil(nb_pages * 0.5))
    return {
        i for i in en_marge
        if _NUMERO_PAGE.match(_cle(elements[i][3])) or (nb_pages >= 2 and len(pages_par_cle[_cle(elements[i][3])]) >= seuil)
    }


# ─── Assemblage ────────────────────────────────────────────────────────────────

_FIN_DE_PHRASE = re.compile(r"[.:;!?»)]$")
# Début de ligne qui ne peut que prolonger la phrase précédente : minuscule, ou
# symbole rejeté à la ligne (« 12.500,00 / € HT », « 3,5 / % »), ponctuation.
_SUITE_DE_PHRASE = re.compile(r"^([a-zà-öø-ÿœæ]|[€%$£)\]»,;:.])")


def fusionner_coupures(blocs: List[Bloc], mots: Set[str]) -> List[Bloc]:
    """Recolle un paragraphe coupé (changement de page, colonne) en pleine phrase."""
    resultat: List[Bloc] = []
    for bloc in blocs:
        precedent = resultat[-1] if resultat else None
        if (
            precedent is not None and precedent.tag == bloc.tag == "p"
            and not _FIN_DE_PHRASE.search(precedent.texte) and _SUITE_DE_PHRASE.match(bloc.texte)
        ):
            texte, html = joindre([(precedent.texte, precedent.html), (bloc.texte, bloc.html)], mots)
            resultat[-1] = Bloc("p", texte, html, precedent.page)
        else:
            resultat.append(bloc)
    return resultat


def html_des_blocs(blocs: Sequence[Bloc]) -> str:
    return "\n".join(f"<{'p' if b.tag in ('p', 'bold') else b.tag}>{b.html}</{'p' if b.tag in ('p', 'bold') else b.tag}>" for b in blocs)


def texte_depuis_html(fragment: str) -> str:
    """Texte d'un HTML : un paragraphe par bloc, séparés par une ligne vide ; une ligne de tableau = « case | case »."""
    fragment = re.sub(r"<br\s*/?>", "\n", fragment, flags=re.I)
    fragment = re.sub(r"</t[dh]>", " | ", fragment, flags=re.I)
    fragment = re.sub(r"</(p|h[1-6]|li|tr|caption|blockquote|div)>", "\x00", fragment, flags=re.I)
    texte = html_module.unescape(re.sub(r"<[^>]+>", "", fragment))
    blocs = (re.sub(r"[ \t]+", " ", re.sub(r"\s*\|\s*$", "", b)).strip() for b in texte.split("\x00"))
    return "\n\n".join(b for b in blocs if b)


# ─── Word ──────────────────────────────────────────────────────────────────────

_STYLES_WORD = "p[style-name='Title'] => h1:fresh\np[style-name='Subtitle'] => h2:fresh"


def lire_word(contenu: bytes) -> DocumentLu:
    if contenu[:4] == b"\xd0\xcf\x11\xe0":  # signature des anciens fichiers binaires (.doc)
        raise ValueError(MESSAGE_DOC)
    if mammoth:
        try:
            fragment = mammoth.convert_to_html(io.BytesIO(contenu), style_map=_STYLES_WORD).value or ""
            # Les images (logos, signatures) ne s'affichent pas dans l'écran d'analyse et
            # alourdissaient la réponse de plusieurs Mo (images intégrées en base64).
            fragment = normaliser(re.sub(r"<img[^>]*>", "", fragment))
            texte = texte_depuis_html(fragment)
            if texte.strip():
                return DocumentLu(texte, fragment, "word", "word", 0)
        except Exception as erreur:
            logger.warning(f"mammoth a échoué : {erreur}")
    if DocxDocument:
        try:
            paragraphes = [normaliser(p.text).strip() for p in DocxDocument(io.BytesIO(contenu)).paragraphs]
            paragraphes = [p for p in paragraphes if p]
            if paragraphes:
                fragment = "".join(f"<p>{html_module.escape(p)}</p>" for p in paragraphes)
                return DocumentLu("\n\n".join(paragraphes), fragment, "word", "word", 0)
        except Exception as erreur:
            logger.warning(f"python-docx a échoué : {erreur}")
    raise ValueError("Impossible de lire ce fichier Word : vérifiez qu'il s'ouvre bien dans Word, puis réessayez.")


# ─── PDF numérique ─────────────────────────────────────────────────────────────

_LISTE_OU_ARTICLE = re.compile(r"^([-•·▪●*]|\(?[a-z0-9]{1,3}[).°]\s|article\s|chapitre\s|titre\s|annexe\s)", re.I)


def _lignes_pdf(doc, pages: Sequence[int]) -> List[dict]:
    """Lignes de texte des pages numériques : html (gras, italique), texte, taille, position."""
    lignes = []
    for numero in pages:
        page = doc[numero]
        hauteur = page.rect.height or 1
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for index, ligne in enumerate(block.get("lines", [])):
                html, texte, tailles, tout_gras = "", "", [], True
                for span in ligne.get("spans", []):
                    brut = normaliser(span.get("text", ""))
                    sur = html_module.escape(brut)
                    if brut.strip():
                        flags = span.get("flags", 0)
                        gras = bool(flags & 16) or "bold" in span.get("font", "").lower()
                        italique = bool(flags & 2)
                        tout_gras = tout_gras and gras
                        tailles.append(span.get("size", 0))
                        if italique:
                            sur = f"<em>{sur}</em>"
                        if gras:
                            sur = f"<strong>{sur}</strong>"
                    html += sur
                    texte += brut
                if not texte.strip():
                    continue
                x0, y0, x1, y1 = ligne.get("bbox", (0, 0, 0, 0))
                lignes.append({
                    "html": html.strip(), "texte": texte.strip(), "taille": max(tailles or [0]), "gras": tout_gras,
                    "x1": x1, "y0": y0, "y1": y1, "nouveau_bloc": index == 0, "page": numero,
                    "haut": y0 / hauteur, "bas": y1 / hauteur,
                })
    return lignes


def blocs_pdf(doc, pages: Sequence[int]) -> List[Bloc]:
    """Paragraphes et titres des pages numériques, lus ligne par ligne comme un lecteur.

    Un bloc PyMuPDF contient souvent toute une page : une ligne en gros caractères ne doit
    pas transformer tout le bloc en titre.
    """
    lignes = _lignes_pdf(doc, pages)
    retirees = parasites([(l["page"], l["haut"], l["bas"], l["texte"]) for l in lignes], len(pages))
    lignes = [l for i, l in enumerate(lignes) if i not in retirees]
    if not lignes:
        return []
    mots = vocabulaire("\n".join(l["texte"] for l in lignes))

    tailles = sorted(l["taille"] for l in lignes)
    taille_mediane = tailles[len(tailles) // 2] or 1
    bord_droit = sorted(l["x1"] for l in lignes)[int(len(lignes) * 0.9)]

    def niveau(ligne: dict) -> str:
        if ligne["taille"] >= taille_mediane * 1.5:
            return "h1"
        if ligne["taille"] >= taille_mediane * 1.2:
            return "h2"
        # Ligne courte entièrement en gras : titre de clause, gardé seul.
        return "bold" if ligne["gras"] and len(ligne["texte"]) <= 100 else "p"

    blocs: List[Bloc] = []
    courant: List[dict] = []
    tag_courant = "p"
    precedente: Optional[dict] = None

    def vider():
        if courant:
            texte, html = joindre([(l["texte"], l["html"]) for l in courant], mots)
            blocs.append(Bloc(tag_courant, texte, html, courant[0]["page"]))
            courant.clear()

    for ligne in lignes:
        tag = niveau(ligne)
        if precedente is None:
            nouveau = True
        elif ligne["page"] != precedente["page"]:
            # Changement de page : le paragraphe continue s'il était coupé en pleine phrase.
            fin_phrase = _FIN_DE_PHRASE.search(precedente["texte"])
            nouveau = not (tag == tag_courant == "p" and not fin_phrase and _SUITE_DE_PHRASE.match(ligne["texte"]))
        else:
            nouveau = (
                tag != tag_courant
                or tag == "bold"
                or ligne["nouveau_bloc"]
                # Espace vertical plus grand qu'un interligne : nouveau paragraphe.
                or ligne["y0"] - precedente["y1"] > (precedente["y1"] - precedente["y0"]) * 0.6
                # Ligne précédente nettement plus courte qui finit une phrase.
                or (_FIN_DE_PHRASE.search(precedente["texte"]) and precedente["x1"] < bord_droit - 40)
                or _LISTE_OU_ARTICLE.match(ligne["texte"]) is not None
            )
        if nouveau:
            vider()
            tag_courant = tag
        courant.append(ligne)
        precedente = ligne
    vider()
    return fusionner_coupures(blocs, mots)


# ─── PDF scanné (Google Vision) ────────────────────────────────────────────────

_LOT_VISION = 5  # pages par appel : au-delà, l'API ignore les pages suivantes
_ESPACE, _ESPACE_SURE, _FIN_LIGNE_ESPACE, _TIRET, _FIN_LIGNE = 1, 2, 3, 4, 5
_TITRE_DE_CLAUSE = re.compile(r"^(article|chapitre|titre|section|annexe|préambule|preambule)\b", re.I)


def _vision_lot(contenu: bytes, lot: Sequence[int]):
    requete = vision.AnnotateFileRequest(
        input_config=vision.InputConfig(content=contenu, mime_type="application/pdf"),
        features=[vision.Feature(type_=vision.Feature.Type.DOCUMENT_TEXT_DETECTION)],
        image_context=vision.ImageContext(language_hints=["fr"]),
        pages=[numero + 1 for numero in lot],
    )
    reponse = vision_client.batch_annotate_files(requests=[requete])
    return list(reponse.responses[0].responses)


def _lignes_paragraphe(paragraphe) -> List[str]:
    """Lignes d'un paragraphe Vision ; un tiret de fin de ligne (absent du texte Vision) est remis pour joindre()."""
    lignes, courante = [], ""
    for mot in paragraphe.words:
        for symbole in mot.symbols:
            courante += symbole.text
            coupure = int(symbole.property.detected_break.type_) if symbole.property else 0
            if coupure in (_ESPACE, _ESPACE_SURE):
                courante += " "
            elif coupure == _TIRET:
                lignes.append(courante + ("" if courante.endswith("-") else "-"))
                courante = ""
            elif coupure in (_FIN_LIGNE_ESPACE, _FIN_LIGNE):
                lignes.append(courante)
                courante = ""
    if courante.strip():
        lignes.append(courante)
    return [espaces_ocr(normaliser(l)).strip() for l in lignes if l.strip()]


def _hauteurs(boite, hauteur_page: float) -> Tuple[float, float]:
    """Haut et bas d'une zone Vision, en fraction de la page (coordonnées normalisées pour les PDF)."""
    if boite.normalized_vertices:
        ys = [v.y for v in boite.normalized_vertices]
    else:
        ys = [v.y / (hauteur_page or 1) for v in boite.vertices]
    return (min(ys), max(ys)) if ys else (0.5, 0.5)


def _tag_ocr(texte: str, premier: bool) -> str:
    court = len(texte) <= 120 and not re.search(r"[.;,]$", texte)
    lettres = [c for c in texte if c.isalpha()]
    majuscules = len(lettres) >= 4 and sum(c.isupper() for c in lettres) >= 0.8 * len(lettres)
    if premier and majuscules and court:
        return "h1"
    if court and (_TITRE_DE_CLAUSE.match(texte) or majuscules):
        return "bold"
    return "p"


def blocs_ocr(contenu: bytes, pages: Sequence[int]) -> List[Bloc]:
    if not (vision and vision_client):
        raise RuntimeError("Lecture des documents scannés indisponible : Google Vision n'est pas configuré.")
    lots = [list(pages[i:i + _LOT_VISION]) for i in range(0, len(pages), _LOT_VISION)]
    with ThreadPoolExecutor(max_workers=min(4, len(lots))) as pool:
        reponses = list(pool.map(lambda lot: _vision_lot(contenu, lot), lots))

    elements = []  # (page, haut, bas, lignes)
    for lot, images in zip(lots, reponses):
        for numero, image in zip(lot, images):
            if image.error.message:
                raise RuntimeError(f"Lecture du scan échouée (page {numero + 1}) : {image.error.message}")
            for page in image.full_text_annotation.pages:
                for block in page.blocks:
                    for paragraphe in block.paragraphs:
                        lignes = _lignes_paragraphe(paragraphe)
                        if lignes:
                            haut, bas = _hauteurs(paragraphe.bounding_box, page.height)
                            elements.append((numero, haut, bas, lignes))

    retires = parasites([(p, h, b, " ".join(l)) for p, h, b, l in elements], len(pages))
    elements = [e for i, e in enumerate(elements) if i not in retires]
    mots = vocabulaire("\n".join(" ".join(l) for *_, l in elements))
    blocs = []
    for numero, _, _, lignes in elements:
        texte, html = joindre([(l, html_module.escape(l)) for l in lignes], mots)
        tag = _tag_ocr(texte, premier=not blocs)
        blocs.append(Bloc(tag, texte, f"<strong>{html}</strong>" if tag == "bold" else html, numero))
    return fusionner_coupures(blocs, mots)


# ─── PDF ───────────────────────────────────────────────────────────────────────

def _page_scannee(page) -> bool:
    """Page sans texte numérique (moins de 30 caractères) mais avec une image : c'est un scan."""
    caracteres, images = 0, 0
    for block in page.get_text("dict").get("blocks", []):
        if block.get("type") == 1:
            images += 1
        elif block.get("type") == 0:
            caracteres += sum(len(s.get("text", "").strip()) for l in block.get("lines", []) for s in l.get("spans", []))
    return caracteres < 30 and images > 0


def lire_pdf(contenu: bytes, forcer_ocr: bool = False) -> DocumentLu:
    if not fitz:
        raise RuntimeError("Lecture des PDF indisponible : PyMuPDF n'est pas installé.")
    try:
        doc = fitz.open(stream=contenu, filetype="pdf")
    except Exception:
        raise ValueError("Ce PDF est illisible ou protégé par un mot de passe.")
    try:
        if doc.needs_pass:
            raise ValueError("Ce PDF est protégé par un mot de passe : retirez la protection, puis réessayez.")
        modes = ["ocr" if forcer_ocr or _page_scannee(page) else "numerique" for page in doc]
        # Pages consécutives de même nature, lues ensemble (en-têtes et paragraphes à cheval).
        series: List[Tuple[str, List[int]]] = []
        for numero, mode in enumerate(modes):
            if series and series[-1][0] == mode:
                series[-1][1].append(numero)
            else:
                series.append((mode, [numero]))
        blocs: List[Bloc] = []
        for mode, pages in series:
            blocs += blocs_ocr(contenu, pages) if mode == "ocr" else blocs_pdf(doc, pages)
        nb_pages = len(doc)
    finally:
        doc.close()

    blocs = fusionner_coupures(blocs, vocabulaire("\n".join(b.texte for b in blocs)))
    if not blocs:
        raise ValueError(MESSAGE_VIDE)
    lecture = modes[0] if len(set(modes)) == 1 else "mixte"
    logger.info(f"PDF lu : {nb_pages} page(s), lecture {lecture}, {len(blocs)} paragraphes.")
    return DocumentLu("\n\n".join(b.texte for b in blocs), html_des_blocs(blocs), "server", lecture, nb_pages)


def lire_document(nom_fichier: str, contenu: bytes, forcer_ocr: bool = False) -> DocumentLu:
    """Point d'entrée : lit un Word (.docx) ou un PDF. ValueError = message à montrer tel quel à l'utilisateur."""
    extension = nom_fichier.rsplit(".", 1)[-1].lower() if "." in nom_fichier else ""
    if extension in ("doc", "docx"):
        return lire_word(contenu)
    return lire_pdf(contenu, forcer_ocr)
