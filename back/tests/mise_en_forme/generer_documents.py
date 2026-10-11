"""
Fabrique les documents de test de la mise en forme (Word, PDF numériques, scans).

Chaque document reprend le même contrat, truffé des cas qui posaient problème :
montants (12.500,00 €), dates et heures (01.02.2026 à 14:30), articles de loi
(art. L.441-10), adresses (www.lumenjuris.com, dpo@lumenjuris.com), noms de
marque (LumenJuris SaaS), mots composés (vice-président, c'est-à-dire), mots
coupés en fin de ligne, en-têtes et pieds de page, listes et tableau.

Chaque article contient un « repère de contrôle n° N » : on vérifie ainsi
qu'aucune page n'est perdue à la lecture (les scans de plus de 5 pages).

Lancement, depuis le dossier lumenjuris :
    back/venv/Scripts/python.exe back/tests/mise_en_forme/generer_documents.py
"""
import io
import random
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from PIL import Image, ImageFilter

DOSSIER = Path(__file__).parent / "documents"

TITRE = "CONTRAT DE PRESTATION DE SERVICES"
EN_TETE = "LumenJuris – Contrat de prestation de services"

PARTIES = [
    "ENTRE LES SOUSSIGNÉS :",
    "La société LumenJuris SaaS, société par actions simplifiée (S.A.S.) au capital de 10.000 €, "
    "immatriculée au RCS de Paris sous le numéro 912 345 678, dont le siège social est situé "
    "12 rue de la Paix, 75002 Paris, représentée par M. Jean Dupont, en sa qualité de vice-président, "
    "ci-après dénommée « le Prestataire »,",
    "ET :",
    "La société Durand Industries, S.A.R.L. au capital de 50.000 €, représentée par Mme Claire Durand, "
    "gérante, ci-après dénommée « le Client ».",
]

TABLEAU = [
    ["Prestation", "Quantité", "Prix unitaire HT"],
    ["Audit juridique initial", "1", "4.500,00 €"],
    ["Accompagnement mensuel", "12", "625,00 €"],
    ["Formation des équipes", "2", "250,00 €"],
]

ARTICLES = [
    ("Objet", [
        "Le présent contrat a pour objet de définir les conditions dans lesquelles le Prestataire fournit "
        "au Client les prestations décrites ci-après, c'est-à-dire l'accompagnement juridique de ses "
        "équipes et la mise à disposition de la plateforme accessible à l'adresse www.lumenjuris.com. "
        "Les prestations sont exécutées conformément aux règles de l'art et aux usages de la profession.",
    ]),
    ("Durée", [
        "Le contrat prend effet le 01.02.2026 à 14:30 pour une durée de douze (12) mois. Il est "
        "renouvelable par tacite reconduction pour des périodes successives d'un an, sauf dénonciation "
        "par lettre recommandée avec accusé de réception adressée au moins trois (3) mois avant "
        "l'échéance contractuelle.",
    ]),
    ("Prix et modalités de paiement", [
        "En contrepartie des prestations, le Client verse au Prestataire la somme forfaitaire de "
        "12.500,00 € HT, soit 15.000,00 € TTC, payable en trois échéances égales. Tout retard de paiement "
        "entraîne de plein droit l'application de pénalités égales à 3,5 % par mois de retard, ainsi qu'une "
        "indemnité forfaitaire pour frais de recouvrement de 40 € conformément à l'art. L.441-10 du Code "
        "de commerce.",
        "TABLEAU",
    ]),
    ("Obligations du Prestataire", [
        "Le Prestataire s'engage à :",
        "a) exécuter les prestations avec diligence et professionnalisme ;",
        "b) informer le Client de toute difficulté susceptible d'affecter l'exécution du contrat ;",
        "c) respecter les délais convenus entre les parties.",
    ]),
    ("Responsabilité", [
        "La responsabilité du Prestataire ne peut être engagée qu'en cas de faute prouvée. En tout état de "
        "cause, l'indemnisation due au Client est irrémédiablement plafonnée au montant des sommes "
        "effectivement perçues au titre des douze derniers mois. Le Prestataire ne saurait être tenu "
        "responsable des conséquences indirectes, notamment les pertes d'exploitation, les manques à gagner "
        "ou les atteintes à l'image, quelles qu'en soient la nature et l'importance.",
    ]),
    ("Confidentialité", [
        "Chacune des parties s'oblige à tenir strictement confidentielles les informations communiquées par "
        "l'autre partie, pendant toute la durée du contrat et cinq (5) ans après son expiration. Cette "
        "obligation ne s'applique pas aux informations tombées dans le domaine public ou dont la "
        "communication est imposée par une disposition légale ou une décision de justice.",
    ]),
    ("Données personnelles", [
        "Les parties s'engagent à respecter le Règlement (UE) 2016/679 du 27 avril 2016 relatif à la "
        "protection des données à caractère personnel. Toute demande relative aux données peut être "
        "adressée à dpo@lumenjuris.com. Le Prestataire agit en qualité de sous-traitant au sens de "
        "l'article 28 du règlement précité.",
    ]),
    ("Résiliation", [
        "En cas de manquement de l'une des parties à ses obligations, non réparé dans un délai de trente "
        "(30) jours à compter de la réception d'une mise en demeure, l'autre partie pourra résilier de plein "
        "droit le présent contrat, sans préjudice de tous dommages et intérêts. Les dispositions de "
        "l'art. L.1242-8 du Code du travail ne sont pas applicables aux présentes.",
    ]),
    ("Droit applicable et juridiction", [
        "Le présent contrat est soumis au droit français. Tout litige relatif à sa formation, son "
        "interprétation ou son exécution relève de la compétence exclusive du Tribunal de commerce de "
        "Paris, nonobstant pluralité de défendeurs ou appel en garantie.",
    ]),
    ("Dispositions diverses", [
        "Le fait pour l'une des parties de ne pas se prévaloir d'un manquement de l'autre partie ne saurait "
        "être interprété comme une renonciation à s'en prévaloir ultérieurement. Le présent contrat "
        "exprime l'intégralité des obligations des parties et annule tout accord antérieur, notamment "
        "la proposition commerciale du 15.12.2025.",
    ]),
]

# Articles supplémentaires pour la version longue (scan de plus de 5 pages).
TITRES_SUPPLEMENTAIRES = [
    "Propriété intellectuelle", "Sous-traitance", "Assurances", "Force majeure", "Non-sollicitation",
    "Cession du contrat", "Notifications", "Indépendance des parties", "Langue du contrat",
    "Élection de domicile", "Révision du prix", "Audit", "Réversibilité", "Garanties",
]
PHRASES = [
    "Chaque partie conserve la propriété des éléments qu'elle apporte à l'exécution du présent contrat.",
    "Le Prestataire ne peut confier tout ou partie des prestations à un tiers sans l'accord écrit du Client.",
    "Le Prestataire déclare être titulaire d'une police d'assurance couvrant sa responsabilité civile professionnelle.",
    "Aucune des parties ne pourra être tenue pour responsable d'un manquement résultant d'un cas de force majeure.",
    "Les notifications sont valablement faites par lettre recommandée avec accusé de réception.",
    "Les parties agissent en toute indépendance et aucune ne peut engager l'autre sans son accord exprès.",
    "En cas de difficulté d'interprétation, la version française du présent contrat prévaut sur toute traduction.",
    "Le prix est révisé chaque année au 1er janvier selon l'évolution de l'indice Syntec publié par l'INSEE.",
    "Le Client peut faire procéder, à ses frais, à un audit de la bonne exécution des prestations.",
    "À l'expiration du contrat, le Prestataire restitue au Client l'ensemble des données dans un format standard.",
]


def articles(longue: bool = False):
    """Liste (titre, paragraphes) ; chaque article porte un repère de contrôle unique."""
    liste = [(titre, list(paragraphes)) for titre, paragraphes in ARTICLES]
    if longue:
        rng = random.Random(7)
        for titre in TITRES_SUPPLEMENTAIRES:
            paragraphe = " ".join(rng.sample(PHRASES, 5))
            liste.append((titre, [paragraphe, " ".join(rng.sample(PHRASES, 4))]))
    for numero, (_, paragraphes) in enumerate(liste, start=1):
        paragraphes.append(f"Repère de contrôle n° {numero}.")
    return liste


# ─── Word ──────────────────────────────────────────────────────────────────────

def creer_word(chemin: Path, titres_avec_styles: bool):
    doc = Document()
    if titres_avec_styles:
        doc.add_heading(TITRE, level=0)
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(TITRE)
        run.bold, run.font.size = True, Pt(16)
    for ligne in PARTIES:
        p = doc.add_paragraph()
        p.add_run(ligne).bold = ligne.endswith(":")
    for numero, (titre, paragraphes) in enumerate(articles(), start=1):
        entete = f"Article {numero} – {titre}"
        if titres_avec_styles:
            doc.add_heading(entete, level=2)
        else:
            p = doc.add_paragraph()
            run = p.add_run(entete)
            run.bold, run.font.size = True, Pt(12)
        for texte in paragraphes:
            if texte == "TABLEAU":
                table = doc.add_table(rows=len(TABLEAU), cols=3)
                table.style = "Table Grid"
                for i, ligne in enumerate(TABLEAU):
                    for j, cellule in enumerate(ligne):
                        table.cell(i, j).text = cellule
            elif titres_avec_styles and texte[:2] in ("a)", "b)", "c)"):
                doc.add_paragraph(texte[3:], style="List Bullet")
            else:
                doc.add_paragraph(texte)
    doc.add_paragraph("Fait à Paris, le 15.01.2026, en deux exemplaires originaux.")
    doc.save(chemin)


# ─── PDF numérique ─────────────────────────────────────────────────────────────

LARGEUR, HAUTEUR, MARGE = 595, 842, 60

# Polices Windows : les polices intégrées de PyMuPDF n'ont ni « € » ni tiret long. La
# plupart des polices Windows se relisent avec des espaces insécables, des tirets
# spéciaux ou des parenthèses décoratives ; Gadugi se relit à l'identique.
POLICES = {"helv": "C:/Windows/Fonts/gadugi.ttf", "hebo": "C:/Windows/Fonts/gadugib.ttf"}
_FONTES = {nom: fitz.Font(fontfile=chemin) for nom, chemin in POLICES.items()}


def _largeur(texte: str, police: str, taille: float) -> float:
    return _FONTES[police].text_length(texte, fontsize=taille)


def _texte(page, position, texte, police="helv", taille=10.5):
    # Nom distinct de « helv » : un nom de police intégrée ferait ignorer le fichier fourni.
    page.insert_text(position, texte, fontname=f"lj{police}", fontfile=POLICES[police], fontsize=taille)


def _decouper_lignes(texte: str, largeur: float, police: str, taille: float, couper_mots: bool):
    """Retour à la ligne simple, avec coupure des mots longs en fin de ligne (comme un PDF justifié)."""
    mesure = lambda s: _largeur(s, police, taille)
    lignes, courante = [], ""
    for mot in texte.split(" "):
        essai = f"{courante} {mot}".strip()
        if mesure(essai) <= largeur:
            courante = essai
            continue
        coupe = None
        if couper_mots and len(mot) >= 9 and "-" not in mot and mot.isalpha():
            for k in range(len(mot) - 3, 3, -1):
                if mesure(f"{courante} {mot[:k]}-".strip()) <= largeur:
                    coupe = k
                    break
        if coupe:
            lignes.append(f"{courante} {mot[:coupe]}-".strip())
            courante = mot[coupe:]
        else:
            lignes.append(courante)
            courante = mot
    if courante:
        lignes.append(courante)
    return lignes


def creer_pdf(chemin: Path, longue=False, colonnes=1, en_tetes=True):
    doc = fitz.open()
    largeur_col = (LARGEUR - 2 * MARGE - (colonnes - 1) * 24) / colonnes
    etat = {"page": None, "y": 0.0, "col": 0}

    def nouvelle_page():
        etat["page"], etat["y"], etat["col"] = doc.new_page(width=LARGEUR, height=HAUTEUR), MARGE + 20, 0

    def place(hauteur):
        if etat["page"] is None:
            nouvelle_page()
        if etat["y"] + hauteur > HAUTEUR - MARGE - 20:
            if etat["col"] + 1 < colonnes:
                etat["col"] += 1
                etat["y"] = MARGE + 20
            else:
                nouvelle_page()
        return MARGE + etat["col"] * (largeur_col + 24)

    def ecrire(texte, police="helv", taille=10.5, avant=0.0, apres=6.0, centre=False):
        etat["y"] += avant
        for ligne in _decouper_lignes(texte, largeur_col, police, taille, couper_mots=police == "helv"):
            x = place(taille * 1.4)
            if centre:
                x += (largeur_col - _largeur(ligne, police, taille)) / 2
            etat["y"] += taille * 1.4
            _texte(etat["page"], (x, etat["y"]), ligne, police, taille)
        etat["y"] += apres

    def tableau():
        largeur_case = largeur_col / 3
        for ligne in TABLEAU:
            x = place(18)
            for j, cellule in enumerate(ligne):
                rect = fitz.Rect(x + j * largeur_case, etat["y"], x + (j + 1) * largeur_case, etat["y"] + 18)
                etat["page"].draw_rect(rect, color=(0, 0, 0), width=0.5)
                _texte(etat["page"], (rect.x0 + 4, rect.y1 - 5), cellule, taille=9)
            etat["y"] += 18
        etat["y"] += 8

    ecrire(TITRE, police="hebo", taille=16, apres=14, centre=True)
    for ligne in PARTIES:
        ecrire(ligne, police="hebo" if ligne.endswith(":") else "helv")
    for numero, (titre, paragraphes) in enumerate(articles(longue), start=1):
        ecrire(f"Article {numero} – {titre}", police="hebo", taille=11.5, avant=8)
        for texte in paragraphes:
            tableau() if texte == "TABLEAU" else ecrire(texte)
    ecrire("Fait à Paris, le 15.01.2026, en deux exemplaires originaux.", avant=10)

    if en_tetes:
        total = len(doc)
        for numero, page in enumerate(doc, start=1):
            _texte(page, (MARGE, 38), EN_TETE, taille=8)
            _texte(page, (LARGEUR / 2 - 20, HAUTEUR - 30), f"Page {numero}/{total}", taille=8)
    doc.save(chemin)
    doc.close()


# ─── Scans (images seules) ─────────────────────────────────────────────────────

def scanner(source: Path, chemin: Path, dpi=200, de_travers=False, numeros_numeriques=False):
    """Transforme un PDF numérique en PDF d'images, comme un scanner (option : page de travers, bruit)."""
    src, out = fitz.open(source), fitz.open()
    rng = random.Random(3)
    for numero, page in enumerate(src, start=1):
        image = Image.open(io.BytesIO(page.get_pixmap(dpi=dpi).tobytes("png"))).convert("L")
        if de_travers:
            image = image.rotate(1.6, expand=False, fillcolor=255, resample=Image.BICUBIC)
            bruit = image.load()
            for _ in range(image.width * image.height // 60):
                x, y = rng.randrange(image.width), rng.randrange(image.height)
                bruit[x, y] = rng.choice((90, 140, 200))
            image = image.filter(ImageFilter.GaussianBlur(0.6))
        tampon = io.BytesIO()
        image.save(tampon, format="JPEG", quality=70 if de_travers else 85)
        nouvelle = out.new_page(width=page.rect.width, height=page.rect.height)
        nouvelle.insert_image(nouvelle.rect, stream=tampon.getvalue())
        if numeros_numeriques:
            # Certains scanners ajoutent un numéro de page en texte : seul texte « numérique » du fichier.
            _texte(nouvelle, (LARGEUR - 90, HAUTEUR - 12), f"Scan {numero}", taille=7)
    out.save(chemin)
    out.close()
    src.close()


def main():
    DOSSIER.mkdir(exist_ok=True)
    creer_word(DOSSIER / "01-word-propre.docx", titres_avec_styles=True)
    creer_word(DOSSIER / "02-word-titres-manuels.docx", titres_avec_styles=False)
    creer_pdf(DOSSIER / "03-pdf-numerique.pdf")
    creer_pdf(DOSSIER / "04-pdf-deux-colonnes.pdf", colonnes=2)
    creer_pdf(DOSSIER / "_source-longue.pdf", longue=True)
    scanner(DOSSIER / "03-pdf-numerique.pdf", DOSSIER / "05-scan-propre.pdf")
    scanner(DOSSIER / "_source-longue.pdf", DOSSIER / "06-scan-long.pdf", dpi=150)
    scanner(DOSSIER / "03-pdf-numerique.pdf", DOSSIER / "07-scan-de-travers.pdf", dpi=150, de_travers=True)
    scanner(DOSSIER / "03-pdf-numerique.pdf", DOSSIER / "08-scan-avec-numeros.pdf", numeros_numeriques=True)
    (DOSSIER / "_source-longue.pdf").unlink()
    # Ancien format Word (.doc, binaire) : seule sa signature compte pour le test du message d'erreur.
    (DOSSIER / "09-ancien-word.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(4096))
    for fichier in sorted(DOSSIER.iterdir()):
        pages = len(fitz.open(fichier)) if fichier.suffix == ".pdf" else "-"
        print(f"{fichier.name:32} {fichier.stat().st_size // 1024:>5} Ko  pages={pages}")


if __name__ == "__main__":
    main()
